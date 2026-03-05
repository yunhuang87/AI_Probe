"""
PPT 生成工具
支持基于模板生成品牌化 PPT，输出存储到本地并可选上传 MinIO
"""
from __future__ import annotations

import tempfile
from pathlib import Path
from typing import Any, Dict, List, Optional
import logging

from pptx import Presentation  # type: ignore

from .storage import upload_to_minio

logger = logging.getLogger(__name__)


CONTENT_TYPE_PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


class PPTGeneratorTool:
    name = "generate_ppt"
    description = "基于模板生成品牌化PPT，支持文本/要点/图片，返回文件路径与下载URL"
    version = "1.0.0"

    async def execute(self, payload: Dict[str, Any]) -> Dict[str, Any]:
        """
        payload:
        - template: 模板路径（可选）
        - sections: [{title, bullets: [str], notes: str}]
        - images: [{path: str, title: str}]
        - brand: {color_primary, color_secondary, font_family} (预留)
        """
        template = payload.get("template")
        sections: List[Dict[str, Any]] = payload.get("sections", []) or []
        images: List[Dict[str, Any]] = payload.get("images", []) or []

        prs = Presentation(template) if template else Presentation()

        # 简单封面
        if sections:
            cover = prs.slides.add_slide(prs.slide_layouts[0])
            cover.shapes.title.text = sections[0].get("title", "自动生成报告")
            if cover.placeholders and len(cover.placeholders) > 1:
                try:
                    cover.placeholders[1].text = payload.get("subtitle", "LuminaOS 自动生成")
                except Exception:
                    pass

        # 章节内容
        for sec in sections:
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = sec.get("title", "")
            body = slide.shapes.placeholders[1].text_frame
            bullets = sec.get("bullets", []) or []
            if not bullets:
                bullets = [sec.get("content", "")]
            for b in bullets:
                body.add_paragraph().text = str(b)
            notes = sec.get("notes")
            if notes:
                slide.notes_slide.notes_text_frame.text = str(notes)

        # 图片占位
        for img in images:
            path = img.get("path")
            if not path:
                continue
            try:
                slide = prs.slides.add_slide(prs.slide_layouts[5])
                slide.shapes.title.text = img.get("title", "图片")
                left = top = width = height = None  # 使用默认占位
                slide.shapes.add_picture(str(path), left or 0, top or 0, width=width, height=height)
            except Exception as e:
                logger.warning(f"Failed to insert image {path}: {e}")

        # 写文件
        tmpdir = Path(tempfile.mkdtemp())
        outfile = tmpdir / "report.pptx"
        prs.save(outfile)

        # 可选上传 MinIO
        download_url: Optional[str] = None
        try:
            download_url = upload_to_minio(outfile, content_type=CONTENT_TYPE_PPTX)
        except Exception as e:
            logger.warning(f"MinIO upload skipped/failure: {e}")

        return {
            "file_path": str(outfile),
            "download_url": download_url,
        }




















