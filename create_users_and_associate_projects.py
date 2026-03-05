#!/usr/bin/env python3
"""
在服务器上批量创建用户账号并关联到项目
"""
import sys
import os
import json
from pathlib import Path

# 尝试导入pandas，如果没有则使用openpyxl（可选，优先使用JSON）
USE_PANDAS = False
USE_OPENPYXL = False
try:
    import pandas as pd
    USE_PANDAS = True
except ImportError:
    try:
        from openpyxl import load_workbook
        USE_OPENPYXL = True
    except ImportError:
        pass  # 如果没有pandas和openpyxl，将使用JSON文件

# 添加项目路径
# 在容器内，路径可能不同
if os.path.exists('/app/src'):
    # 容器内路径
    sys.path.insert(0, '/app/src')
    sys.path.insert(0, '/app')
    # 尝试添加可能的database路径
    for db_path in ['/app/database/src', '/app/../database/src', '/app/../../database/src']:
        if os.path.exists(db_path):
            sys.path.insert(0, db_path)
else:
    # 本地路径
    project_root = Path(__file__).parent
    sys.path.insert(0, str(project_root / 'auth-service' / 'src'))
    sys.path.insert(0, str(project_root / 'database' / 'src'))
    sys.path.insert(0, str(project_root / 'project-management' / 'src'))

from sqlalchemy.orm import Session
from passlib.context import CryptContext
from uuid import UUID
from datetime import date

# 尝试导入数据库相关模块
try:
    from database.src.core.database import get_db_session
    from database.src.repositories.user_repository import UserRepository
    from database.src.models.user_models import UserStatus
    # 延迟导入项目模型，避免初始化问题
    PROJECT_MODELS_AVAILABLE = True
except ImportError:
    # 尝试从auth-service的导入方式
    try:
        from src.core.database import SessionLocal, init_database
        from src.repositories.user_repository import UserRepository
        from database.src.models.user_models import UserStatus
        PROJECT_MODELS_AVAILABLE = True
        
        def get_db_session():
            init_database()
            return SessionLocal()
    except ImportError:
        print("Error: Cannot import database modules")
        import traceback
        traceback.print_exc()
        sys.exit(1)

# 配置
# 支持多个可能的路径
EXCEL_FILE = "中化国际数字化部员工统计_20251204.xlsx"
JSON_FILE = "users_from_excel.json"  # 如果Excel读取失败，使用JSON文件
DEFAULT_PASSWORD = "Sinochem@2025"
DEFAULT_ROLE = "member"  # 默认项目角色：manager, member, viewer

# 可能的文件路径
POSSIBLE_PATHS = [
    "/tmp",
    "/app",
    ".",
]

# 密码加密
pwd_context = CryptContext(schemes=['bcrypt'], deprecated='auto')

def read_users_from_json(json_file: str):
    """从JSON文件读取用户信息"""
    try:
        with open(json_file, 'r', encoding='utf-8') as f:
            users = json.load(f)
        return users
    except Exception as e:
        print(f"Error reading JSON file: {str(e)}")
        return []

def read_excel_users(excel_file: str):
    """读取Excel文件中的用户信息"""
    try:
        if USE_PANDAS:
            df = pd.read_excel(excel_file, sheet_name=0)
            
            # 识别列
            name_col = None
            email_col = None
            
            for col in df.columns:
                col_str = str(col)
                if '姓名' in col_str or 'name' in col_str.lower():
                    name_col = col
                elif '邮箱' in col_str or 'email' in col_str.lower() or 'mail' in col_str.lower():
                    email_col = col
            
            if not name_col or not email_col:
                print(f"❌ 无法识别姓名或邮箱列。可用列: {list(df.columns)}")
                return []
            
            # 提取用户数据
            users = []
            for idx, row in df.iterrows():
                name = str(row[name_col]).strip() if pd.notna(row[name_col]) else ""
                email = str(row[email_col]).strip() if pd.notna(row[email_col]) else ""
                
                if not email:
                    continue
                
                # 从邮箱提取用户名
                username = email.split('@')[0] if '@' in email else email
                
                users.append({
                    'name': name,
                    'email': email,
                    'username': username
                })
            
            return users
        else:
            # 使用openpyxl
            wb = load_workbook(excel_file, read_only=True)
            ws = wb.active
            
            # 读取表头
            headers = [cell.value for cell in ws[1]]
            
            # 识别列索引
            name_idx = None
            email_idx = None
            
            for idx, header in enumerate(headers):
                header_str = str(header) if header else ""
                if '姓名' in header_str or 'name' in header_str.lower():
                    name_idx = idx
                elif '邮箱' in header_str or 'email' in header_str.lower() or 'mail' in header_str.lower():
                    email_idx = idx
            
            if name_idx is None or email_idx is None:
                print(f"❌ 无法识别姓名或邮箱列。可用列: {headers}")
                return []
            
            # 读取数据
            users = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                name = str(row[name_idx]).strip() if row[name_idx] else ""
                email = str(row[email_idx]).strip() if row[email_idx] else ""
                
                if not email or email == 'None':
                    continue
                
                # 从邮箱提取用户名
                username = email.split('@')[0] if '@' in email else email
                
                users.append({
                    'name': name,
                    'email': email,
                    'username': username
                })
            
            return users
    except Exception as e:
        print(f"❌ 读取Excel文件失败: {str(e)}")
        import traceback
        traceback.print_exc()
        return []

def create_user(user_info: dict, db: Session) -> tuple:
    """创建用户，返回 (success, user_id, message)"""
    try:
        user_repo = UserRepository(db)
        
        # 检查用户是否已存在
        existing = user_repo.get_by_username(user_info['username'])
        if existing:
            return (True, str(existing.id), f"用户已存在: {user_info['username']}")
        
        # 检查邮箱是否已存在
        existing_email = user_repo.get_by_email(user_info['email'])
        if existing_email:
            return (True, str(existing_email.id), f"邮箱已存在: {user_info['email']}")
        
        # 加密密码
        password_hash = pwd_context.hash(DEFAULT_PASSWORD)
        
        # 创建用户
        user = user_repo.create_user(
            username=user_info['username'],
            email=user_info['email'],
            password_hash=password_hash,
            full_name=user_info['name'],
            status=UserStatus.ACTIVE
        )
        
        db.commit()
        return (True, str(user.id), f"用户创建成功: {user_info['username']}")
        
    except Exception as e:
        db.rollback()
        return (False, None, f"创建失败: {str(e)}")

def get_all_projects(db: Session):
    """获取所有项目"""
    try:
        # 延迟导入，避免模型初始化问题
        from database.src.models.project_models import Project
        projects = db.query(Project).order_by(Project.created_at.desc()).all()
        return projects
    except Exception as e:
        print(f"Error getting project list: {str(e)}")
        # 如果模型导入失败，尝试直接查询表
        try:
            from sqlalchemy import text
            result = db.execute(text("SELECT id, name, project_code FROM pm_projects ORDER BY created_at DESC"))
            projects = []
            for row in result:
                # 创建简单的项目对象
                class SimpleProject:
                    def __init__(self, id, name, project_code):
                        self.id = id
                        self.name = name
                        self.project_code = project_code
                projects.append(SimpleProject(row[0], row[1], row[2]))
            return projects
        except Exception as e2:
            print(f"Error querying projects directly: {str(e2)}")
            return []

def add_user_to_project(user_id: str, project_id: str, role: str, db: Session) -> tuple:
    """添加用户到项目，返回 (success, message)"""
    try:
        # 延迟导入，避免模型初始化问题
        from database.src.models.project_models import ProjectMember, ProjectMemberRole
        
        user_uuid = UUID(user_id)
        project_uuid = UUID(project_id)
        
        # 检查用户是否已是项目成员
        existing = db.query(ProjectMember).filter(
            ProjectMember.project_id == project_uuid,
            ProjectMember.user_id == user_uuid
        ).first()
        
        if existing:
            return (True, f"用户已是项目成员（角色: {existing.role}）")
        
        # 确定角色
        if role == "manager":
            member_role = ProjectMemberRole.MANAGER.value
        elif role == "viewer":
            member_role = ProjectMemberRole.VIEWER.value
        else:
            member_role = ProjectMemberRole.MEMBER.value
        
        # 创建项目成员
        member = ProjectMember(
            project_id=project_uuid,
            user_id=user_uuid,
            role=member_role,
            joined_at=date.today()
        )
        
        db.add(member)
        db.commit()
        return (True, f"已添加为项目成员（角色: {role}）")
        
    except Exception as e:
        db.rollback()
        # 如果模型导入失败，尝试直接SQL插入
        try:
            from sqlalchemy import text
            user_uuid = UUID(user_id)
            project_uuid = UUID(project_id)
            
            # 确定角色值
            if role == "manager":
                member_role = "manager"
            elif role == "viewer":
                member_role = "viewer"
            else:
                member_role = "member"
            
            # 检查是否已存在
            check_sql = text("""
                SELECT id FROM pm_project_members 
                WHERE project_id = :project_id AND user_id = :user_id
            """)
            result = db.execute(check_sql, {"project_id": project_uuid, "user_id": user_uuid}).first()
            if result:
                return (True, f"用户已是项目成员")
            
            # 插入新成员
            insert_sql = text("""
                INSERT INTO pm_project_members (id, project_id, user_id, role, joined_at, created_at, updated_at)
                VALUES (gen_random_uuid(), :project_id, :user_id, :role, :joined_at, NOW(), NOW())
            """)
            db.execute(insert_sql, {
                "project_id": project_uuid,
                "user_id": user_uuid,
                "role": member_role,
                "joined_at": date.today()
            })
            db.commit()
            return (True, f"已添加为项目成员（角色: {role}）")
        except Exception as e2:
            import traceback
            traceback.print_exc()
            return (False, f"添加失败: {str(e2)}")

def main():
    """主函数"""
    print("=" * 80)
    print("批量创建用户账号并关联到项目")
    print("=" * 80)
    print("")
    
    # 读取用户信息（优先使用JSON，如果没有则使用Excel）
    users = []
    json_path = None
    excel_path = None
    
    # 尝试在不同路径查找文件
    for base_path in POSSIBLE_PATHS:
        test_json = os.path.join(base_path, JSON_FILE) if base_path != "." else JSON_FILE
        test_excel = os.path.join(base_path, EXCEL_FILE) if base_path != "." else EXCEL_FILE
        
        if os.path.exists(test_json):
            json_path = test_json
            break
        if os.path.exists(test_excel):
            excel_path = test_excel
    
    if json_path:
        print(f"Reading users from JSON file: {json_path}")
        users = read_users_from_json(json_path)
    elif excel_path and (USE_PANDAS or USE_OPENPYXL):
        print(f"Reading Excel file: {excel_path}")
        users = read_excel_users(excel_path)
    else:
        print(f"Error: Cannot find {JSON_FILE} or {EXCEL_FILE}")
        print(f"Searched in: {POSSIBLE_PATHS}")
        print("Please upload the JSON file or Excel file to the server")
        sys.exit(1)
    
    if not users:
        print("Error: No user data found")
        sys.exit(1)
    
    print(f"✅ 读取成功，共 {len(users)} 个用户")
    print("")
    
    # 获取数据库会话
    db = get_db_session()
    
    try:
        # 步骤1: 创建用户
        print("=" * 80)
        print("步骤1: 创建用户账号")
        print("=" * 80)
        print("")
        
        user_map = {}  # username -> user_id
        success_count = 0
        skip_count = 0
        fail_count = 0
        
        for user_info in users:
            print(f"处理用户: {user_info['name']} ({user_info['username']})")
            success, user_id, message = create_user(user_info, db)
            
            if success:
                user_map[user_info['username']] = {
                    'user_id': user_id,
                    'name': user_info['name'],
                    'email': user_info['email']
                }
                if "已存在" in message:
                    skip_count += 1
                    print(f"  ⚠️  {message}")
                else:
                    success_count += 1
                    print(f"  ✅ {message}")
            else:
                fail_count += 1
                print(f"  ❌ {message}")
        
        print("")
        print(f"创建统计: 成功 {success_count}, 跳过 {skip_count}, 失败 {fail_count}")
        print("")
        
        # 步骤2: 获取项目列表
        print("=" * 80)
        print("步骤2: 获取项目列表")
        print("=" * 80)
        print("")
        
        projects = get_all_projects(db)
        
        if not projects:
            print("⚠️  未找到项目，跳过项目关联")
            return
        
        print(f"✅ 找到 {len(projects)} 个项目:")
        for idx, project in enumerate(projects, 1):
            print(f"  {idx}. {project.name} (ID: {project.id})")
        print("")
        
        # 选择项目（默认关联所有项目）
        # 支持通过环境变量配置，否则使用默认值
        project_selection = os.getenv('PROJECT_SELECTION', 'all').strip()
        role_selection = os.getenv('PROJECT_ROLE', '2').strip()
        
        if project_selection.lower() == 'skip':
            print("✅ 跳过项目关联（环境变量设置）")
            return
        
        if not project_selection or project_selection.lower() == 'all':
            selected_projects = projects
            print(f"✅ 已选择所有 {len(projects)} 个项目")
        else:
            try:
                indices = [int(x.strip()) - 1 for x in project_selection.split(',')]
                selected_projects = [projects[i] for i in indices if 0 <= i < len(projects)]
                print(f"✅ 已选择 {len(selected_projects)} 个项目")
            except:
                print(f"⚠️  无效的项目选择: {project_selection}，使用所有项目")
                selected_projects = projects
        
        print("")
        
        # 选择角色
        if role_selection == "1":
            role = "manager"
        elif role_selection == "3":
            role = "viewer"
        else:
            role = "member"
        
        print(f"✅ 已选择角色: {role} (可通过环境变量 PROJECT_ROLE 修改)")
        print("")
        
        # 步骤3: 关联用户到项目
        print("=" * 80)
        print("步骤3: 关联用户到项目")
        print("=" * 80)
        print("")
        
        total_associations = 0
        success_associations = 0
        skip_associations = 0
        fail_associations = 0
        
        for project in selected_projects:
            print(f"项目: {project.name}")
            for username, user_data in user_map.items():
                if not user_data['user_id']:
                    continue
                
                total_associations += 1
                success, message = add_user_to_project(
                    user_data['user_id'],
                    str(project.id),
                    role,
                    db
                )
                
                if success:
                    if "已是" in message:
                        skip_associations += 1
                        print(f"  ⚠️  {user_data['name']}: {message}")
                    else:
                        success_associations += 1
                        print(f"  ✅ {user_data['name']}: {message}")
                else:
                    fail_associations += 1
                    print(f"  ❌ {user_data['name']}: {message}")
            
            print("")
        
        print("=" * 80)
        print("关联统计")
        print("=" * 80)
        print(f"✅ 成功: {success_associations}")
        print(f"⚠️  跳过: {skip_associations}")
        print(f"❌ 失败: {fail_associations}")
        print(f"📊 总计: {total_associations}")
        print("")
        
        # 保存用户映射
        with open('user_id_map.json', 'w', encoding='utf-8') as f:
            json.dump(user_map, f, ensure_ascii=False, indent=2)
        print("✅ 用户ID映射已保存到: user_id_map.json")
        
    finally:
        db.close()
    
    print("")
    print("=" * 80)
    print("完成！")
    print("=" * 80)
    print(f"默认密码: {DEFAULT_PASSWORD}")
    print("提示: 用户首次登录后应修改密码")

if __name__ == "__main__":
    main()

