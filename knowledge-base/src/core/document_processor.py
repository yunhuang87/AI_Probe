"""
文档处理器
支持多种文档格式的解析和内容提取
"""
import logging
import os
from pathlib import Path
from typing import List, Dict, Any, Optional, Tuple
from datetime import datetime

from ..models.document_models import (
    DocumentType, DocumentChunk, ChunkMetadata, DocumentMetadata
)
from ..config import settings

logger = logging.getLogger(__name__)

# 尝试导入文档处理库
try:
    import PyPDF2
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False
    logger.warning("PyPDF2 not available")

try:
    from docx import Document as DocxDocument
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    logger.warning("python-docx not available")

try:
    import openpyxl
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False
    logger.warning("openpyxl not available")


class DocumentProcessor:
    """文档处理器"""
    
    def __init__(self, chunk_size: int = None, chunk_overlap: int = None):
        """
        初始化文档处理器
        
        Args:
            chunk_size: 分块大小
            chunk_overlap: 分块重叠大小
        """
        self.chunk_size = chunk_size or settings.CHUNK_SIZE
        self.chunk_overlap = chunk_overlap or settings.CHUNK_OVERLAP
    
    def detect_file_type(self, file_path: str) -> DocumentType:
        """
        检测文件类型
        
        Args:
            file_path: 文件路径
        
        Returns:
            文档类型
        """
        ext = Path(file_path).suffix.lower()
        
        type_mapping = {
            '.pdf': DocumentType.PDF,
            '.doc': DocumentType.WORD,
            '.docx': DocumentType.WORD,
            '.xls': DocumentType.EXCEL,
            '.xlsx': DocumentType.EXCEL,
            '.txt': DocumentType.TEXT,
            '.md': DocumentType.MARKDOWN,
            '.markdown': DocumentType.MARKDOWN,
        }
        
        return type_mapping.get(ext, DocumentType.UNKNOWN)
    
    def extract_text(self, file_path: str, file_type: Optional[DocumentType] = None) -> Tuple[str, DocumentMetadata]:
        """
        提取文档文本内容
        
        Args:
            file_path: 文件路径
            file_type: 文件类型（可选，自动检测）
        
        Returns:
            (文本内容, 元数据)
        """
        if file_type is None:
            file_type = self.detect_file_type(file_path)
        
        metadata = DocumentMetadata()
        
        try:
            if file_type == DocumentType.PDF:
                return self._extract_from_pdf(file_path, metadata)
            elif file_type == DocumentType.WORD:
                return self._extract_from_word(file_path, metadata)
            elif file_type == DocumentType.EXCEL:
                return self._extract_from_excel(file_path, metadata)
            elif file_type == DocumentType.TEXT:
                return self._extract_from_text(file_path, metadata)
            elif file_type == DocumentType.MARKDOWN:
                return self._extract_from_text(file_path, metadata)
            else:
                raise ValueError(f"Unsupported file type: {file_type}")
        except Exception as e:
            logger.error(f"Error extracting text from {file_path}: {str(e)}")
            raise
    
    def _extract_from_pdf(self, file_path: str, metadata: DocumentMetadata) -> Tuple[str, DocumentMetadata]:
        """从PDF提取文本"""
        if not PDF_AVAILABLE:
            raise ImportError("PyPDF2 is required for PDF processing")
        
        text_parts = []
        
        with open(file_path, 'rb') as file:
            pdf_reader = PyPDF2.PdfReader(file)
            metadata.page_count = len(pdf_reader.pages)
            
            for page_num, page in enumerate(pdf_reader.pages, 1):
                text = page.extract_text()
                text_parts.append(text)
            
            # 尝试提取元数据
            if pdf_reader.metadata:
                metadata.title = pdf_reader.metadata.get('/Title')
                metadata.author = pdf_reader.metadata.get('/Author')
                if pdf_reader.metadata.get('/CreationDate'):
                    try:
                        metadata.creation_date = pdf_reader.metadata.get('/CreationDate')
                    except:
                        pass
        
        content = '\n\n'.join(text_parts)
        metadata.word_count = len(content.split())
        
        return content, metadata
    
    def _extract_from_word(self, file_path: str, metadata: DocumentMetadata) -> Tuple[str, DocumentMetadata]:
        """从Word文档提取文本"""
        if not DOCX_AVAILABLE:
            raise ImportError("python-docx is required for Word processing")
        
        doc = DocxDocument(file_path)
        
        # 提取文本
        paragraphs = [para.text for para in doc.paragraphs]
        content = '\n\n'.join(paragraphs)
        
        # 提取元数据
        if doc.core_properties:
            metadata.title = doc.core_properties.title
            metadata.author = doc.core_properties.author
            if doc.core_properties.created:
                metadata.creation_date = doc.core_properties.created
            if doc.core_properties.modified:
                metadata.modification_date = doc.core_properties.modified
        
        metadata.word_count = len(content.split())
        
        return content, metadata
    
    def _extract_from_excel(self, file_path: str, metadata: DocumentMetadata) -> Tuple[str, DocumentMetadata]:
        """从Excel提取文本"""
        if not EXCEL_AVAILABLE:
            raise ImportError("openpyxl is required for Excel processing")
        
        workbook = openpyxl.load_workbook(file_path, data_only=True)
        
        text_parts = []
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            sheet_text = []
            for row in sheet.iter_rows(values_only=True):
                row_text = ' | '.join(str(cell) if cell is not None else '' for cell in row)
                sheet_text.append(row_text)
            text_parts.append(f"Sheet: {sheet_name}\n" + '\n'.join(sheet_text))
        
        content = '\n\n'.join(text_parts)
        metadata.word_count = len(content.split())
        
        return content, metadata
    
    def _extract_from_text(self, file_path: str, metadata: DocumentMetadata) -> Tuple[str, DocumentMetadata]:
        """从文本文件提取"""
        with open(file_path, 'r', encoding='utf-8') as file:
            content = file.read()
        
        metadata.word_count = len(content.split())
        
        return content, metadata
    
    def chunk_text(
        self,
        text: str,
        metadata: Optional[Dict[str, Any]] = None,
        strategy: str = "fixed"
    ) -> List[DocumentChunk]:
        """
        将文本分块
        
        Args:
            text: 文本内容
            metadata: 额外元数据
            strategy: 分块策略 (fixed, semantic, hierarchical, table, code, smart)
        
        Returns:
            文档块列表
        """
        from .chunking_strategies import ChunkingStrategyFactory
        
        # 使用智能分块策略
        chunking_strategy = ChunkingStrategyFactory.create_strategy(
            strategy=strategy,
            chunk_size=self.chunk_size,
            chunk_overlap=self.chunk_overlap
        )
        
        return chunking_strategy.chunk(text, metadata)
    
    def process_document(
        self,
        file_path: str,
        chunking_strategy: str = "smart"
    ) -> Tuple[str, DocumentMetadata, List[DocumentChunk]]:
        """
        处理文档：提取文本并分块
        
        Args:
            file_path: 文件路径
            chunking_strategy: 分块策略
        
        Returns:
            (文本内容, 元数据, 文档块列表)
        """
        text, metadata = self.extract_text(file_path)
        chunks = self.chunk_text(text, strategy=chunking_strategy)
        
        return text, metadata, chunks


# 全局文档处理器实例
_document_processor: Optional[DocumentProcessor] = None


def get_document_processor() -> DocumentProcessor:
    """获取文档处理器单例"""
    global _document_processor
    
    if _document_processor is None:
        _document_processor = DocumentProcessor()
    
    return _document_processor









