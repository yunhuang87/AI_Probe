"""
分析项目管理Excel表格，生成项目管理数据
"""
import sys
import os
import json
import uuid
from pathlib import Path
from typing import Dict, List, Any, Optional
from datetime import datetime, date
import re

# 添加项目路径
sys.path.insert(0, str(Path(__file__).parent))
sys.path.insert(0, str(Path(__file__).parent / "database"))
sys.path.insert(0, str(Path(__file__).parent / "shared_libs"))

try:
    import openpyxl
    from openpyxl import load_workbook
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False
    print("警告: openpyxl未安装，请运行: pip install openpyxl")

import psycopg2
from psycopg2.extras import execute_values

DB_CONFIG = {
    'host': 'localhost',
    'port': 5432,
    'database': 'ai_platform',
    'user': 'ai_user',
    'password': 'ai_password'
}

class ProjectExcelAnalyzer:
    """项目管理Excel分析器"""
    
    def __init__(self, excel_path: str):
        self.excel_path = excel_path
        self.wb = None
        self.projects = []
        self.phases = []
        self.milestones = []
        self.tasks = []
        
    def load_excel(self):
        """加载Excel文件"""
        if not EXCEL_AVAILABLE:
            raise ImportError("openpyxl未安装，请运行: pip install openpyxl")
        
        print(f"加载Excel文件: {self.excel_path}")
        self.wb = load_workbook(self.excel_path, data_only=True)
        print(f"工作表: {self.wb.sheetnames}")
    
    def analyze_sheet(self, sheet_name: Optional[str] = None):
        """分析工作表"""
        if sheet_name:
            ws = self.wb[sheet_name]
        else:
            ws = self.wb.active
        
        print(f"\n分析工作表: {ws.title}")
        print(f"行数: {ws.max_row}, 列数: {ws.max_column}")
        
        # 读取表头（通常在第2行）
        headers = []
        header_row = 2
        header_row_data = list(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))[0]
        headers = [str(cell).strip() if cell else '' for cell in header_row_data]
        
        print(f"表头行: {header_row}")
        print(f"表头: {headers[:8]}")
        
        # 分析数据结构
        projects_data = []
        current_project = None
        project_index = 0
        
        for row_idx, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
            if row_idx > ws.max_row:
                break
            
            # 获取项目名称列（通常是第2列，索引1）
            project_name_cell = row[1] if len(row) > 1 else None
            project_name = str(project_name_cell).strip() if project_name_cell else ''
            
            # 跳过空行或表头行
            if not project_name or project_name in ['项目名称', '']:
                continue
            
            # 检查是否是新的项目行（项目名称不为空）
            if project_name and len(project_name) > 1 and project_name not in ['计划', '成果']:
                # 保存上一个项目
                if current_project:
                    projects_data.append(current_project)
                
                # 创建新项目
                project_index += 1
                current_project = {
                    'name': project_name,
                    'project_code': f"PRJ-{project_index:04d}",
                    'tasks': [],
                    'metadata': {}
                }
                
                # 提取项目信息
                if len(row) > 2 and row[2]:  # 填报人
                    current_project['metadata']['reporter'] = str(row[2]).strip()
                
                if len(row) > 3 and row[3]:  # 项目状态
                    status_str = str(row[3]).strip()
                    current_project['status'] = self._map_status(status_str)
                    current_project['metadata']['original_status'] = status_str
                
                if len(row) > 4 and row[4]:  # 项目阶段
                    current_project['metadata']['phase'] = str(row[4]).strip()
                
                if len(row) > 5 and row[5]:  # 待办及重点提示
                    current_project['description'] = str(row[5]).strip()
                    # 将待办事项作为任务
                    todo_text = str(row[5]).strip()
                    if todo_text:
                        current_project['tasks'].append({
                            'name': f"待办事项: {todo_text[:50]}",
                            'description': todo_text,
                            'status': 'todo'
                        })
                
                # 提取周报信息（从第7列开始）
                weekly_tasks = []
                for col_idx in range(6, min(len(row), len(headers))):
                    if row[col_idx] and headers[col_idx]:
                        cell_value = str(row[col_idx]).strip()
                        if cell_value and cell_value not in ['计划', '成果']:
                            weekly_tasks.append({
                                'week': headers[col_idx],
                                'content': cell_value
                            })
                
                if weekly_tasks:
                    current_project['metadata']['weekly_reports'] = weekly_tasks
                    # 将周报内容作为任务
                    for wt in weekly_tasks:
                        current_project['tasks'].append({
                            'name': f"{wt['week']}: {wt['content'][:50]}",
                            'description': wt['content'],
                            'status': 'in_progress'
                        })
            elif current_project:
                # 可能是项目的续行（计划/成果）
                # 提取周报内容
                for col_idx in range(6, min(len(row), len(headers))):
                    if row[col_idx] and headers[col_idx]:
                        cell_value = str(row[col_idx]).strip()
                        if cell_value:
                            # 添加到周报或任务
                            week_name = headers[col_idx]
                            row_type = str(row[6]).strip() if len(row) > 6 else ''  # 计划/成果
                            
                            task_name = f"{week_name} - {row_type}: {cell_value[:50]}"
                            current_project['tasks'].append({
                                'name': task_name,
                                'description': cell_value,
                                'status': 'completed' if row_type == '成果' else 'in_progress'
                            })
        
        # 保存最后一个项目
        if current_project:
            projects_data.append(current_project)
        
        return projects_data
    
    def _extract_project_info(self, row_data: Dict, headers: List[str]) -> Optional[Dict]:
        """提取项目信息"""
        project_info = {}
        
        # 查找项目相关字段
        project_fields = {
            '项目': ['项目名称', '项目', '项目代码', '项目编码', '项目编号'],
            'code': ['项目代码', '项目编码', '项目编号', '编码', '代码'],
            'status': ['状态', '项目状态', '进度状态'],
            'start_date': ['开始日期', '计划开始', '启动日期'],
            'end_date': ['结束日期', '计划结束', '完成日期', '截止日期'],
            'progress': ['进度', '完成度', '进度百分比', '%'],
            'manager': ['负责人', '项目经理', '负责人姓名'],
            'priority': ['优先级', '重要程度'],
            'budget': ['预算', '项目预算', '预算金额'],
        }
        
        # 检查是否是项目行（通常项目名称在特定列）
        project_name = None
        for field in project_fields['项目']:
            for header in headers:
                if field in header and row_data.get(header):
                    project_name = str(row_data[header]).strip()
                    break
            if project_name:
                break
        
        if not project_name or len(project_name) < 2:
            return None
        
        project_info['name'] = project_name
        project_info['is_new_project'] = True
        
        # 提取项目编码
        project_code = None
        for field in project_fields['code']:
            for header in headers:
                if field in header and row_data.get(header):
                    project_code = str(row_data[header]).strip()
                    break
            if project_code:
                break
        
        if not project_code:
            # 生成项目编码
            project_code = f"PRJ-{project_name[:10].upper().replace(' ', '-')}"
        
        project_info['project_code'] = project_code
        
        # 提取其他字段
        for key, fields in project_fields.items():
            if key == '项目' or key == 'code':
                continue
            for field in fields:
                for header in headers:
                    if field in header and row_data.get(header):
                        value = row_data[header]
                        if value:
                            project_info[key] = self._parse_value(value, key)
                            break
                if key in project_info:
                    break
        
        # 设置默认值
        if 'status' in project_info:
            project_info['status'] = self._map_status(project_info['status'])
        else:
            project_info.setdefault('status', 'active')
        project_info.setdefault('priority', 'medium')
        project_info.setdefault('progress_percent', 0.0)
        
        return project_info
    
    def _extract_task_info(self, row_data: Dict, headers: List[str]) -> Optional[Dict]:
        """提取任务信息"""
        task_info = {}
        
        # 查找任务相关字段
        task_fields = {
            'name': ['任务', '工作项', '活动', '事项'],
            'status': ['状态', '任务状态'],
            'assignee': ['负责人', '执行人', '分配人'],
            'start_date': ['开始日期', '计划开始'],
            'due_date': ['截止日期', '计划完成', '结束日期'],
            'progress': ['进度', '完成度', '进度百分比'],
            'estimated_hours': ['预估工时', '预计工时', '工时'],
        }
        
        task_name = None
        for field in task_fields['name']:
            for header in headers:
                if field in header and row_data.get(header):
                    task_name = str(row_data[header]).strip()
                    break
            if task_name:
                break
        
        if not task_name or len(task_name) < 2:
            return None
        
        task_info['name'] = task_name
        
        # 提取其他字段
        for key, fields in task_fields.items():
            if key == 'name':
                continue
            for field in fields:
                for header in headers:
                    if field in header and row_data.get(header):
                        value = row_data[header]
                        if value:
                            task_info[key] = self._parse_value(value, key)
                            break
                if key in task_info:
                    break
        
        # 设置默认值
        task_info.setdefault('status', 'todo')
        task_info.setdefault('progress_percent', 0.0)
        
        return task_info
    
    def _map_status(self, status_str: str) -> str:
        """映射状态字符串到枚举值"""
        status_str = status_str.lower().strip()
        
        # 项目状态映射
        status_map = {
            'planning': 'planning',
            'active': 'active',
            'in_progress': 'active',
            '正常进行': 'active',
            '02-正常进行': 'active',
            '01-准备/可研': 'planning',
            '准备': 'planning',
            '可研': 'planning',
            'completed': 'completed',
            '完成': 'completed',
            'delayed': 'delayed',
            '延迟': 'delayed',
            'cancelled': 'cancelled',
            '取消': 'cancelled',
        }
        
        for key, value in status_map.items():
            if key in status_str:
                return value
        
        return 'active'  # 默认值
    
    def _parse_value(self, value: Any, field_type: str) -> Any:
        """解析值"""
        if value is None:
            return None
        
        if field_type in ['start_date', 'end_date', 'due_date', 'target_date']:
            # 解析日期
            if isinstance(value, datetime):
                return value.date()
            elif isinstance(value, date):
                return value
            elif isinstance(value, str):
                # 尝试解析日期字符串
                try:
                    return datetime.strptime(value, '%Y-%m-%d').date()
                except:
                    try:
                        return datetime.strptime(value, '%Y/%m/%d').date()
                    except:
                        return None
            return None
        elif field_type in ['progress', 'progress_percent']:
            # 解析进度百分比
            if isinstance(value, (int, float)):
                return float(value)
            elif isinstance(value, str):
                # 提取数字
                match = re.search(r'(\d+\.?\d*)', value.replace('%', ''))
                if match:
                    return float(match.group(1))
            return 0.0
        elif field_type in ['budget', 'actual_cost', 'estimated_hours', 'actual_hours']:
            # 解析数字
            if isinstance(value, (int, float)):
                return float(value)
            elif isinstance(value, str):
                # 提取数字
                match = re.search(r'(\d+\.?\d*)', value.replace(',', ''))
                if match:
                    return float(match.group(1))
            return None
        else:
            return str(value).strip() if value else None
    
    def analyze_all_sheets(self):
        """分析所有工作表"""
        all_projects = []
        
        for sheet_name in self.wb.sheetnames:
            try:
                projects = self.analyze_sheet(sheet_name)
                all_projects.extend(projects)
            except Exception as e:
                print(f"分析工作表 {sheet_name} 时出错: {e}")
                continue
        
        return all_projects


class ProjectDataGenerator:
    """项目管理数据生成器"""
    
    def __init__(self, db_config: Dict[str, str]):
        self.db_config = db_config
        self.conn = None
    
    def connect(self):
        """连接数据库"""
        self.conn = psycopg2.connect(**self.db_config)
    
    def close(self):
        """关闭连接"""
        if self.conn:
            self.conn.close()
    
    def insert_projects(self, projects: List[Dict[str, Any]]):
        """插入项目"""
        if not projects:
            return []
        
        cur = self.conn.cursor()
        project_ids = []
        
        try:
            for project in projects:
                project_code = project.get('project_code', f"PRJ-{uuid.uuid4().hex[:8].upper()}")
                
                # 先检查项目是否已存在
                cur.execute("SELECT id FROM pm_projects WHERE project_code = %s", (project_code,))
                existing = cur.fetchone()
                
                if existing:
                    project_id = str(existing[0])
                    # 更新现有项目
                    cur.execute("""
                        UPDATE pm_projects 
                        SET name = %s,
                            description = %s,
                            status = %s,
                            progress_percent = %s,
                            updated_at = now()
                        WHERE id = %s
                    """, (
                        project['name'],
                        project.get('description', ''),
                        project.get('status', 'active'),
                        project.get('progress_percent', 0.0),
                        project_id
                    ))
                else:
                    # 插入新项目
                    project_id = str(uuid.uuid4())
                    cur.execute("""
                        INSERT INTO pm_projects 
                        (id, project_code, name, description, status, priority, start_date, end_date, 
                         progress_percent, budget, metadata)
                        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                    """, (
                        project_id,
                        project_code,
                        project['name'],
                        project.get('description', ''),
                        project.get('status', 'active'),
                        project.get('priority', 'medium'),
                        project.get('start_date'),
                        project.get('end_date'),
                        project.get('progress_percent', 0.0),
                        project.get('budget'),
                        json.dumps(project.get('metadata', {}), ensure_ascii=False)
                    ))
                
                project_ids.append(project_id)
            
            self.conn.commit()
            print(f"   [OK] 插入 {len(projects)} 个项目")
            return project_ids
        except Exception as e:
            self.conn.rollback()
            print(f"   [ERROR] 插入项目失败: {e}")
            import traceback
            traceback.print_exc()
            return []
        finally:
            cur.close()
    
    def insert_phases(self, project_id: str, phases: List[Dict[str, Any]]):
        """插入项目阶段"""
        if not phases:
            return []
        
        cur = self.conn.cursor()
        phase_ids = []
        
        try:
            for idx, phase in enumerate(phases):
                phase_id = str(uuid.uuid4())
                phase_ids.append(phase_id)
                
                cur.execute("""
                    INSERT INTO pm_project_phases 
                    (id, project_id, name, description, sequence, start_date, end_date, progress_percent, metadata)
                    VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
                """, (
                    phase_id,
                    project_id,
                    phase.get('name', f'阶段{idx + 1}'),
                    phase.get('description', ''),
                    phase.get('sequence', idx),
                    phase.get('start_date'),
                    phase.get('end_date'),
                    phase.get('progress_percent', 0.0),
                    json.dumps(phase.get('metadata', {}), ensure_ascii=False)
                ))
            
            self.conn.commit()
            if phases:
                print(f"   [OK] 插入 {len(phases)} 个项目阶段")
            return phase_ids
        except Exception as e:
            self.conn.rollback()
            print(f"   [ERROR] 插入项目阶段失败: {e}")
            return []
        finally:
            cur.close()
    
    def insert_tasks(self, project_id: str, phase_id: Optional[str], tasks: List[Dict[str, Any]]):
        """插入任务"""
        if not tasks:
            return []
        
        cur = self.conn.cursor()
        task_ids = []
        
        try:
            for task in tasks:
                task_id = str(uuid.uuid4())
                task_ids.append(task_id)
                
                cur.execute("""
                    INSERT INTO pm_tasks 
                    (id, project_id, phase_id, name, description, status, assignee_id, 
                     start_date, due_date, progress_percent, estimated_hours, metadata)
                    VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                """, (
                    task_id,
                    project_id,
                    phase_id,
                    task.get('name', ''),
                    task.get('description', ''),
                    task.get('status', 'todo'),
                    task.get('assignee_id'),
                    task.get('start_date'),
                    task.get('due_date'),
                    task.get('progress_percent', 0.0),
                    task.get('estimated_hours'),
                    json.dumps(task.get('metadata', {}), ensure_ascii=False)
                ))
            
            self.conn.commit()
            if tasks:
                print(f"   [OK] 插入 {len(tasks)} 个任务")
            return task_ids
        except Exception as e:
            self.conn.rollback()
            print(f"   [ERROR] 插入任务失败: {e}")
            return []
        finally:
            cur.close()


def main():
    """主函数"""
    excel_path = "2025项目周月进度报告 (1).xlsx"
    
    if not os.path.exists(excel_path):
        print(f"错误: 文件不存在: {excel_path}")
        return
    
    try:
        # 分析Excel
        print("=" * 60)
        print("分析项目管理Excel表格")
        print("=" * 60)
        
        analyzer = ProjectExcelAnalyzer(excel_path)
        analyzer.load_excel()
        projects_data = analyzer.analyze_all_sheets()
        
        if not projects_data:
            print("\n未找到项目数据，尝试直接读取所有行...")
            # 尝试更简单的方法
            ws = analyzer.wb.active
            projects_data = []
            current_project = None
            
            for row_idx, row in enumerate(ws.iter_rows(min_row=1, values_only=True), 1):
                if row_idx > 100:  # 限制行数
                    break
                
                # 查找包含项目信息的行
                row_str = ' '.join([str(cell) for cell in row if cell])
                if '项目' in row_str or len(row_str) > 20:
                    if not current_project or '项目' in row_str:
                        if current_project:
                            projects_data.append(current_project)
                        current_project = {
                            'name': row_str[:100],
                            'project_code': f"PRJ-{row_idx:04d}",
                            'status': 'active',
                            'tasks': []
                        }
                    elif current_project:
                        current_project['tasks'].append({
                            'name': row_str[:100],
                            'status': 'todo'
                        })
            
            if current_project:
                projects_data.append(current_project)
        
        print(f"\n找到 {len(projects_data)} 个项目")
        for p in projects_data:
            print(f"  - {p.get('name', '未知项目')}: {len(p.get('tasks', []))} 个任务")
        
        # 生成数据
        print("\n" + "=" * 60)
        print("生成项目管理数据")
        print("=" * 60)
        
        generator = ProjectDataGenerator(DB_CONFIG)
        generator.connect()
        
        try:
            print("\n插入数据到数据库...")
            total_projects = 0
            total_tasks = 0
            
            for project_data in projects_data:
                # 插入项目
                project_ids = generator.insert_projects([project_data])
                if project_ids:
                    project_id = project_ids[0]
                    total_projects += 1
                    
                    # 插入阶段（如果有）
                    phases = project_data.get('phases', [])
                    phase_ids = generator.insert_phases(project_id, phases)
                    
                    # 插入任务
                    tasks = project_data.get('tasks', [])
                    phase_id = phase_ids[0] if phase_ids else None
                    generator.insert_tasks(project_id, phase_id, tasks)
                    total_tasks += len(tasks)
            
            print("\n" + "=" * 60)
            print("完成！")
            print("=" * 60)
            print(f"\n生成的数据:")
            print(f"  - 项目: {total_projects} 个")
            print(f"  - 任务: {total_tasks} 个")
            
        finally:
            generator.close()
    
    except Exception as e:
        print(f"\n错误: {e}")
        import traceback
        traceback.print_exc()


if __name__ == '__main__':
    main()

